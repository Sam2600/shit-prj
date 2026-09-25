"""Import an .xlsx or .csv file into a new MySQL table billy.{parameter}_tbl.

Usage:
    python import_data.py Sample.xlsx "store master"
    python import_data.py users.csv "user master" --encoding cp932
    python import_data.py Sample.xlsx "store master" --sheet Sheet2 --dry-run
"""

import argparse
import configparser
import csv
import datetime
import os
import re
import sys
from pathlib import Path

VARCHAR_LEN = 512
BATCH_SIZE = 1000
MAX_IDENTIFIER_LEN = 64  # MySQL limit for table / column names
# InnoDB row limit is 65535 bytes; VARCHAR(512) utf8mb4 = 512*4 + 2 bytes each.
MAX_COLUMNS = (65535 - 4) // (VARCHAR_LEN * 4 + 2)
ROW_COLUMN = "Row"
CHARSET = "utf8mb4"
COLLATION = "utf8mb4_unicode_ci"
CONFIG_FILE = Path(__file__).with_name("config.ini")
ENV_FILE = Path(__file__).with_name(".env")


class ImportError_(Exception):
    """A user-facing error: printed without a traceback."""


# ---------------------------------------------------------------- reading

def cell_to_str(value):
    if value is None:
        return None
    if isinstance(value, datetime.datetime):
        if value.time() == datetime.time(0, 0):
            return value.strftime("%Y-%m-%d")
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, (datetime.date, datetime.time)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return text if text != "" else None


def read_xlsx(path, sheet_name):
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name:
            if sheet_name not in wb.sheetnames:
                raise ImportError_(
                    f"Sheet '{sheet_name}' not found. Sheets: {', '.join(wb.sheetnames)}")
            ws = wb[sheet_name]
        else:
            ws = wb.worksheets[0]
        return [[cell_to_str(v) for v in row] for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def read_csv(path, encoding):
    encodings = [encoding] if encoding else ["utf-8-sig", "cp932"]
    for enc in encodings:
        try:
            with open(path, newline="", encoding=enc) as f:
                rows = [[cell_to_str(v) for v in row] for row in csv.reader(f)]
            if not encoding and enc != "utf-8-sig":
                print(f"Note: file is not UTF-8, read it as {enc}.")
            return rows
        except UnicodeDecodeError:
            continue
    raise ImportError_(f"Cannot decode CSV with {', '.join(encodings)}. Try --encoding.")


def read_file(path, sheet_name, encoding):
    ext = path.suffix.lower()
    if ext == ".xlsx":
        return read_xlsx(path, sheet_name)
    if ext == ".csv":
        return read_csv(path, encoding)
    raise ImportError_(f"Unsupported file type '{ext}'. Use .xlsx or .csv.")


# ------------------------------------------------------------- preparing

def table_name_from_param(param):
    name = re.sub(r"[^0-9a-zA-Z]+", "_", param.strip()).strip("_").lower()
    if not name:
        raise ImportError_(f"Parameter '{param}' gives an empty table name.")
    name += "_tbl"
    if len(name) > MAX_IDENTIFIER_LEN:
        raise ImportError_(f"Table name '{name}' is longer than {MAX_IDENTIFIER_LEN} chars.")
    return name


def prepare(rows):
    """Return (columns, data_rows). Every non-empty row is data (no header row);
    columns are named col_1..col_N and rows are padded to the widest row."""
    # Keep (source row number, row) so errors can point at the right line.
    numbered = [(i + 1, r) for i, r in enumerate(rows) if any(v is not None for v in r)]
    if not numbered:
        raise ImportError_("File is empty.")

    def last_filled(row):
        return max((i + 1 for i, v in enumerate(row) if v is not None), default=0)

    width = max(last_filled(r) for _, r in numbered)
    if width > MAX_COLUMNS:
        raise ImportError_(
            f"File has {width} columns; VARCHAR({VARCHAR_LEN}) allows at most "
            f"{MAX_COLUMNS} per MySQL row.")

    columns = [f"col_{i + 1}" for i in range(width)]
    data = []
    for line_no, row in numbered:
        row = (list(row) + [None] * width)[:width]
        for col, value in zip(columns, row):
            if value is not None and len(value) > VARCHAR_LEN:
                raise ImportError_(
                    f"Row {line_no}, column '{col}': value is {len(value)} chars "
                    f"(max {VARCHAR_LEN}).")
        data.append(row)
    return columns, data


def quote_ident(name):
    return "`" + name.replace("`", "``") + "`"


def create_table_sql(database, table, columns):
    col_defs = ",\n  ".join(f"{quote_ident(c)} VARCHAR({VARCHAR_LEN}) NULL" for c in columns)
    return (
        f"CREATE TABLE {quote_ident(database)}.{quote_ident(table)} (\n"
        f"  {quote_ident(ROW_COLUMN)} INT NOT NULL AUTO_INCREMENT PRIMARY KEY,\n"
        f"  {col_defs}\n"
        # Explicit collation: MySQL 8's default (utf8mb4_0900_ai_ci) doesn't exist
        # on MySQL 5.7 / MariaDB, so this keeps the table portable across servers.
        f") ENGINE=InnoDB DEFAULT CHARSET={CHARSET} COLLATE={COLLATION}"
    )


# -------------------------------------------------------------- database

def read_env_file():
    """Parse simple KEY=VALUE lines from .env next to this script."""
    values = {}
    if ENV_FILE.exists():
        for line in ENV_FILE.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if line and not line.startswith("#") and "=" in line:
                key, value = line.split("=", 1)
                values[key.strip()] = value.strip().strip('"').strip("'")
    return values


def load_config():
    """Each setting: environment variable > .env file > config.ini > default."""
    ini = configparser.ConfigParser()
    ini.read(CONFIG_FILE, encoding="utf-8")
    section = ini["mysql"] if ini.has_section("mysql") else {}
    env_file = read_env_file()

    def get(key, env_key, default=None):
        for source in (os.environ, env_file):
            if source.get(env_key):
                return source[env_key]
        return section.get(key, default)

    cfg = {
        "host": get("host", "DB_HOST", "localhost"),
        "port": int(get("port", "DB_PORT", 3306)),
        "user": get("user", "DB_USER"),
        "password": get("password", "DB_PASSWORD", ""),
        "database": get("database", "DB_NAME"),
    }
    for key, env_key in (("database", "DB_NAME"), ("user", "DB_USER")):
        if not cfg[key]:
            raise ImportError_(
                f"{key} not set. Set {env_key} in .env or {key} in config.ini "
                f"(copy .env.example or config.ini.example).")
    return cfg


def write_to_mysql(cfg, table, columns, data, create_sql):
    import mysql.connector

    conn = mysql.connector.connect(
        host=cfg["host"], port=cfg["port"], user=cfg["user"],
        password=cfg["password"], database=cfg["database"],
        charset=CHARSET, collation=COLLATION, autocommit=False,
    )
    full_name = f"{quote_ident(cfg['database'])}.{quote_ident(table)}"
    try:
        cur = conn.cursor()
        cur.execute(f"DROP TABLE IF EXISTS {full_name}")
        cur.execute(create_sql)
        insert_sql = (
            f"INSERT INTO {full_name} ({', '.join(quote_ident(c) for c in columns)}) "
            f"VALUES ({', '.join(['%s'] * len(columns))})"
        )
        try:
            for start in range(0, len(data), BATCH_SIZE):
                cur.executemany(insert_sql, data[start:start + BATCH_SIZE])
            conn.commit()
        except Exception:
            conn.rollback()
            cur.execute(f"DROP TABLE IF EXISTS {full_name}")  # don't leave a half-loaded table
            raise
    finally:
        conn.close()


# ------------------------------------------------------------------ main

def main():
    parser = argparse.ArgumentParser(description="Import .xlsx/.csv into MySQL table {parameter}_tbl")
    parser.add_argument("file", type=Path, help="path to .xlsx or .csv file")
    parser.add_argument("parameter", help='table name parameter, e.g. "store master"')
    parser.add_argument("--sheet", help="Excel sheet name (default: first sheet)")
    parser.add_argument("--encoding", help="CSV encoding (default: UTF-8, fallback cp932)")
    parser.add_argument("--dry-run", action="store_true",
                        help="show the table SQL and row count without touching the DB")
    args = parser.parse_args()

    try:
        if not args.file.is_file():
            raise ImportError_(f"File not found: {args.file}")
        table = table_name_from_param(args.parameter)
        columns, data = prepare(read_file(args.file, args.sheet, args.encoding))

        cfg = load_config()
        if args.dry_run:
            print(create_table_sql(cfg["database"], table, columns) + ";")
            print(f"\n{len(data)} rows would be inserted into {cfg['database']}.{table}.")
            return 0

        write_to_mysql(cfg, table, columns, data, create_table_sql(cfg["database"], table, columns))
        print(f"Created {cfg['database']}.{table} ({len(columns)} columns) "
              f"and inserted {len(data)} rows.")
        return 0
    except ImportError_ as e:
        print(f"Error: {e}", file=sys.stderr)
        return 1
    except Exception as e:  # DB/connection errors
        print(f"Error: {type(e).__name__}: {e}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
